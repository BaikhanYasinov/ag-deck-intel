"""Извлечение содержимого дека.

PDF обрабатывается полностью: текстовый слой плюс растр каждой страницы.
PPTX: если в системе есть LibreOffice, файл конвертируется в PDF и идёт
по полному пути. Если нет — деградированный режим: текст, встроенные
изображения и данные нативных диаграмм, без растра слайда целиком.
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import fitz  # PyMuPDF

from ..config import MAX_PAGES, RENDER_DPI, STORAGE_DIR


@dataclass
class PageData:
    index: int
    text: str = ""
    image_path: str | None = None
    ocr_used: bool = False


@dataclass
class ExtractResult:
    pages: list[PageData] = field(default_factory=list)
    pdf_path: str | None = None
    visual_mode: str = "full"          # full | degraded
    warnings: list[str] = field(default_factory=list)


def _soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    return None


def _convert_to_pdf(src: Path) -> Path | None:
    exe = _soffice()
    if not exe:
        return None
    out_dir = Path(tempfile.mkdtemp())
    try:
        subprocess.run(
            [exe, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(src)],
            check=True, timeout=180, capture_output=True,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
        return None
    pdfs = list(out_dir.glob("*.pdf"))
    return pdfs[0] if pdfs else None


def _ocr(image_path: str) -> str:
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(image_path), lang="rus+eng").strip()
    except Exception:
        return ""


def _extract_pdf(pdf_path: Path, doc_id: str, result: ExtractResult) -> None:
    pages_dir = STORAGE_DIR / "pages" / doc_id
    pages_dir.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(pdf_path)
    if doc.page_count > MAX_PAGES:
        result.warnings.append(
            f"В документе {doc.page_count} страниц, обработаны первые {MAX_PAGES}")
    for i in range(min(doc.page_count, MAX_PAGES)):
        page = doc[i]
        num = i + 1
        img_path = pages_dir / f"{num:03d}.png"
        page.get_pixmap(dpi=RENDER_DPI).save(img_path)
        text = page.get_text().strip()
        ocr_used = False
        if len(text) < 20:
            recovered = _ocr(str(img_path))
            if recovered:
                text, ocr_used = recovered, True
        result.pages.append(PageData(num, text, str(img_path), ocr_used))
    doc.close()


def _extract_pptx_degraded(src: Path, doc_id: str, result: ExtractResult) -> None:
    from pptx import Presentation

    pages_dir = STORAGE_DIR / "pages" / doc_id
    pages_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(src))

    for i, slide in enumerate(prs.slides, start=1):
        if i > MAX_PAGES:
            break
        chunks: list[str] = []
        first_image: str | None = None

        for n, shape in enumerate(slide.shapes):
            if shape.has_text_frame and shape.text_frame.text.strip():
                chunks.append(shape.text_frame.text.strip())
            if getattr(shape, "image", None) is not None:
                try:
                    img = shape.image
                    p = pages_dir / f"{i:03d}_{n}.{img.ext}"
                    p.write_bytes(img.blob)
                    first_image = first_image or str(p)
                except Exception:
                    pass
            if getattr(shape, "has_chart", False):
                try:
                    plot = shape.chart.plots[0]
                    cats = [str(c) for c in plot.categories]
                    for series in plot.series:
                        pairs = ", ".join(
                            f"{c}: {v}" for c, v in zip(cats, series.values) if v is not None)
                        chunks.append(f"[диаграмма] {series.name}: {pairs}")
                except Exception:
                    chunks.append("[диаграмма: данные не прочитаны]")

        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                chunks.append("[заметки] " + slide.notes_slide.notes_text_frame.text.strip())
        except Exception:
            pass

        result.pages.append(PageData(i, "\n".join(chunks), first_image))

    result.visual_mode = "degraded"
    result.warnings.append(
        "LibreOffice не найден: слайды не отрисованы целиком. Визуальный анализ "
        "ограничен текстом, встроенными картинками и данными диаграмм."
    )


def extract_document(src_path: str, doc_id: str) -> ExtractResult:
    src = Path(src_path)
    result = ExtractResult()
    ext = src.suffix.lower()

    if ext == ".pdf":
        _extract_pdf(src, doc_id, result)
        result.pdf_path = str(src)
        return result

    if ext == ".pptx":
        converted = _convert_to_pdf(src)
        if converted:
            target = STORAGE_DIR / "pdf" / f"{doc_id}.pdf"
            shutil.copy(converted, target)
            _extract_pdf(target, doc_id, result)
            result.pdf_path = str(target)
        else:
            _extract_pptx_degraded(src, doc_id, result)
        return result

    raise ValueError(f"Формат {ext} не поддерживается. Допустимы .pdf и .pptx")
